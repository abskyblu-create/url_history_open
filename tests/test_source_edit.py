import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from source_edit import delete_from_source, remove_url_text


class SourceEditTests(unittest.TestCase):
    def test_exact_url_only(self):
        self.assertEqual(remove_url_text('https://example.com/a https://example.com/ab', 'https://example.com/a'), ' https://example.com/ab')
        self.assertEqual(remove_url_text('See www.example.com.', 'https://www.example.com'), 'See .')

    def test_csv_removes_all_occurrences_and_preserves_other_cells_and_backup(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'source.csv'
            source.write_text('name;url\nA;https://example.com/a\nB;"https://example.com/a https://example.com/b"\n', encoding='utf-8-sig')
            original = source.read_bytes()
            changed, backup = delete_from_source(source, 'https://example.com/a')
            self.assertEqual(changed, 2)
            self.assertEqual(backup.read_bytes(), original)
            with source.open(encoding='utf-8-sig', newline='') as stream:
                rows = list(csv.reader(stream, delimiter=';'))
            self.assertEqual(rows, [['name', 'url'], ['A', ''], ['B', ' https://example.com/b']])

    def test_excel_links_formulas_and_other_data(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'source.xlsx'
            book = Workbook()
            sheet = book.active
            sheet['A1'] = 'https://example.com/a'
            sheet['B1'] = 'Keep me'
            sheet['B1'].font = __import__('openpyxl').styles.Font(bold=True)
            sheet['C1'] = '=SUM(1,2)'
            sheet['A2'] = 'Link label'
            sheet['A2'].hyperlink = 'https://example.com/a'
            other = book.create_sheet('More')
            other['A1'] = '=HYPERLINK("https://example.com/a","Example")'
            other['B1'] = 'https://example.com/b'
            book.save(source)
            book.close()
            changed, backup = delete_from_source(source, 'https://example.com/a')
            self.assertEqual(changed, 3)
            self.assertTrue(backup.exists())
            book = load_workbook(source)
            try:
                self.assertIsNone(book.active['A1'].value)
                self.assertEqual(book.active['A2'].value, 'Link label')
                self.assertIsNone(book.active['A2'].hyperlink)
                self.assertEqual(book.active['B1'].value, 'Keep me')
                self.assertTrue(book.active['B1'].font.bold)
                self.assertEqual(book.active['C1'].value, '=SUM(1,2)')
                self.assertIsNone(book['More']['A1'].value)
                self.assertEqual(book['More']['B1'].value, 'https://example.com/b')
            finally:
                book.close()

    def test_failed_save_preserves_source_and_backup(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'source.csv'
            source.write_text('https://example.com/a')
            original = source.read_bytes()
            with patch('source_edit.os.replace', side_effect=PermissionError('File open in Excel')):
                with self.assertRaises(PermissionError):
                    delete_from_source(source, 'https://example.com/a')
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(len(list(Path(directory).glob('*.bak'))), 1)

    def test_missing_url_does_not_rewrite_source(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'source.csv'
            source.write_text('https://example.com/b')
            original = source.read_bytes()
            with self.assertRaises(ValueError):
                delete_from_source(source, 'https://example.com/a')
            self.assertEqual(source.read_bytes(), original)


if __name__ == '__main__':
    unittest.main()
